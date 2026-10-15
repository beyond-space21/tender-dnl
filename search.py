"""Typesense index of extracted tender packet files."""

import hashlib
import json
import os
import re
import time
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import quote

ROOT = Path(__file__).resolve().parent
TENDERS = ROOT / "TENDERS"
COLLECTION = "tender_files"
TEXT_LIMIT = 50_000
TEXT_EXTENSIONS = {".txt", ".csv", ".html", ".htm", ".xml", ".md"}
INDEX_EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".docx", ".xlsx", ".xls", ".doc"}


def _load_env() -> None:
    import db

    db._load_env()


def client():
    import typesense

    _load_env()
    key = os.environ.get("TYPESENSE_API_KEY", "").strip()
    if not key:
        raise RuntimeError("Add TYPESENSE_API_KEY to .env.")
    host = os.environ.get("TYPESENSE_HOST", "127.0.0.1").strip() or "127.0.0.1"
    port = os.environ.get("TYPESENSE_PORT", "8108").strip() or "8108"
    return typesense.Client(
        {
            "nodes": [{"host": host, "port": port, "protocol": "http"}],
            "api_key": key,
            "connection_timeout_seconds": 8,
        }
    )


SYNONYM_SETS = {
    "boq": ["boq", "bill of quantities", "bill of quantity", "price bid", "item rate"],
    "emd": ["emd", "earnest money", "earnest money deposit", "bid security"],
    "nit": ["nit", "notice inviting tender", "tender notice", "invitation to tender"],
    "gst": ["gst", "goods and services tax", "gstin"],
    "turnover": ["turnover", "annual turnover", "average annual turnover", "aat"],
    "dpr": ["dpr", "detailed project report"],
    "loa": ["loa", "letter of acceptance", "letter of award"],
    "work_order": ["work order", "workorder", "purchase order", "po"],
}


def ensure_collection(connection=None) -> None:
    connection = connection or client()
    schema = {
        "name": COLLECTION,
        "fields": [
            {"name": "tender_id", "type": "string"},
            {"name": "title", "type": "string", "infix": True},
            {"name": "organisation", "type": "string", "optional": True, "infix": True},
            {"name": "filename", "type": "string", "infix": True},
            {"name": "relative_path", "type": "string"},
            {"name": "extension", "type": "string", "facet": True},
            {"name": "content", "type": "string"},
            {"name": "closing_at", "type": "int64"},
        ],
        "token_separators": ["-", "_", "/", "."],
    }
    try:
        connection.collections[COLLECTION].retrieve()
    except Exception as exc:
        if "404" not in str(exc) and "Not Found" not in str(exc):
            raise
        connection.collections.create(schema)
    ensure_synonyms(connection)


def ensure_synonyms(connection=None) -> None:
    connection = connection or client()
    for name, words in SYNONYM_SETS.items():
        connection.collections[COLLECTION].synonyms.upsert(
            name, {"synonyms": words}
        )


def _supports_infix(connection) -> bool:
    info = connection.collections[COLLECTION].retrieve()
    fields = {field["name"]: field for field in info.get("fields") or []}
    return bool(fields.get("filename", {}).get("infix"))


def _cap(text: str) -> str:
    cleaned = re.sub(r"\s+", " ", text).strip()
    return cleaned[:TEXT_LIMIT]


def _plain_text(path: Path) -> str:
    raw = path.read_bytes().decode("utf-8", "replace")
    if path.suffix.lower() in {".html", ".htm"}:
        raw = re.sub(r"(?is)<script\b.*?</script>", " ", raw)
        raw = re.sub(r"(?is)<style\b.*?</style>", " ", raw)
        raw = re.sub(r"(?s)<[^>]+>", " ", raw)
    return _cap(raw)


def _pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts: list[str] = []
    size = 0
    for page in reader.pages:
        piece = page.extract_text() or ""
        parts.append(piece)
        size += len(piece)
        if size >= TEXT_LIMIT:
            break
    return _cap("\n".join(parts))


def _docx_text(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
            if sum(len(part) for part in parts) >= TEXT_LIMIT:
                return _cap("\n".join(parts))
    return _cap("\n".join(parts))


def _xlsx_text(path: Path) -> str:
    from openpyxl import load_workbook

    book = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    size = 0
    try:
        for sheet in book.worksheets:
            for row in sheet.iter_rows(values_only=True):
                line = " ".join("" if cell is None else str(cell) for cell in row).strip()
                if not line:
                    continue
                parts.append(line)
                size += len(line)
                if size >= TEXT_LIMIT:
                    return _cap("\n".join(parts))
    finally:
        book.close()
    return _cap("\n".join(parts))


def _xls_text(path: Path) -> str:
    import xlrd

    book = xlrd.open_workbook(str(path))
    parts: list[str] = []
    size = 0
    for sheet in book.sheets():
        for row_index in range(sheet.nrows):
            values = []
            for column in range(sheet.ncols):
                value = sheet.cell_value(row_index, column)
                if value != "":
                    values.append(str(value))
            line = " ".join(values).strip()
            if not line:
                continue
            parts.append(line)
            size += len(line)
            if size >= TEXT_LIMIT:
                return _cap("\n".join(parts))
    return _cap("\n".join(parts))


def extract_text(path: Path) -> str | None:
    extension = path.suffix.lower()
    if extension not in INDEX_EXTENSIONS:
        return None
    if extension == ".doc":
        return ""
    try:
        if extension in TEXT_EXTENSIONS:
            return _plain_text(path)
        if extension == ".pdf":
            return _pdf_text(path)
        if extension == ".docx":
            return _docx_text(path)
        if extension == ".xlsx":
            return _xlsx_text(path)
        if extension == ".xls":
            return _xls_text(path)
    except Exception as exc:
        print(f"unreadable {path.name}: {exc}", flush=True)
        return ""
    return None


def _closing_unix(value) -> int:
    if value is None or value == "":
        return 0
    if isinstance(value, datetime):
        moment = value if value.tzinfo else value.replace(tzinfo=timezone.utc)
        return int(moment.timestamp())
    if isinstance(value, str):
        moment = datetime.fromisoformat(value)
        if moment.tzinfo is None:
            moment = moment.replace(tzinfo=timezone.utc)
        return int(moment.timestamp())
    return 0


def _document_id(tender_id: str, relative_path: str) -> str:
    digest = hashlib.sha1(f"{tender_id}\n{relative_path}".encode()).hexdigest()
    return digest


def _upsert(connection, documents: list[dict]) -> None:
    if not documents:
        return
    raw = connection.collections[COLLECTION].documents.import_(
        documents, {"action": "upsert"}
    )
    lines = raw.splitlines() if isinstance(raw, str) else list(raw)
    for line in lines:
        item = line if isinstance(line, dict) else json.loads(line)
        if not item.get("success", True):
            raise RuntimeError(item.get("error") or "Typesense rejected a document.")


def index_tender(tender: dict, stop_event=None) -> int:
    import db

    tender_id = tender["tender_id"]
    folder = (TENDERS / tender_id).resolve()
    if not folder.is_dir() or TENDERS.resolve() not in folder.parents:
        return 0
    started = time.monotonic()
    connection = client()
    ensure_collection(connection)
    batch: list[dict] = []
    indexed = 0
    for path in sorted(folder.rglob("*")):
        if stop_event is not None and stop_event.is_set():
            break
        if not path.is_file():
            continue
        content = extract_text(path)
        if content is None:
            continue
        relative = path.relative_to(folder).as_posix()
        batch.append(
            {
                "id": _document_id(tender_id, relative),
                "tender_id": tender_id,
                "title": tender.get("title") or "",
                "organisation": tender.get("organisation") or "",
                "filename": path.name,
                "relative_path": relative,
                "extension": path.suffix.lower().lstrip("."),
                "content": content,
                "closing_at": _closing_unix(tender.get("closing_at")),
            }
        )
        if len(batch) >= 20:
            _upsert(connection, batch)
            indexed += len(batch)
            batch = []
    if batch:
        _upsert(connection, batch)
        indexed += len(batch)
    db.merge_timings(
        tender_id,
        {
            "index": {
                "ms": int(round((time.monotonic() - started) * 1000)),
                "files": indexed,
            }
        },
    )
    return indexed


def lifecycle_from_unix(closing_at: int) -> str:
    if not closing_at:
        return "undated"
    now = datetime.now(timezone.utc).timestamp()
    if closing_at <= now:
        return "expired"
    if closing_at <= now + 72 * 3600:
        return "closing_soon"
    return "open"


def _snippet_html(hit: dict) -> str:
    for item in hit.get("highlights") or []:
        raw = item.get("snippet") or ""
        cleaned = re.sub(r"\s+", " ", raw).strip()
        if cleaned:
            return cleaned[:400]
    content = (hit.get("document") or {}).get("content") or ""
    return re.sub(r"\s+", " ", content).strip()[:280]


def _matched_fields(hit: dict) -> list[str]:
    names: list[str] = []
    for item in hit.get("highlights") or []:
        field = item.get("field")
        if field and field not in names:
            names.append(field)
    return names


def _build_filter(
    include_expired: bool,
    lifecycle: str | None,
    extensions: list[str],
) -> str:
    now = int(datetime.now(timezone.utc).timestamp())
    clauses: list[str] = []
    if lifecycle == "expired":
        clauses.append(f"closing_at:<{now}")
    elif lifecycle == "closing_soon":
        clauses.append(f"closing_at:>={now}")
        clauses.append(f"closing_at:<={now + 72 * 3600}")
    elif lifecycle == "open":
        clauses.append(f"closing_at:>{now + 72 * 3600}")
    elif not include_expired:
        clauses.append(f"closing_at:>={now}")
    if extensions:
        quoted = ",".join(extensions)
        clauses.append(f"extension:=[{quoted}]")
    return " && ".join(clauses)


def _facet_rows(result: dict, field: str) -> list[dict]:
    for block in result.get("facet_counts") or []:
        if block.get("field_name") != field:
            continue
        rows = []
        for item in block.get("counts") or []:
            rows.append({"value": item.get("value") or "", "count": item.get("count") or 0})
        return rows
    return []


def _count_map(values: list[str]) -> list[dict]:
    tallies: dict[str, int] = {}
    for value in values:
        key = value or "unknown"
        tallies[key] = tallies.get(key, 0) + 1
    return [
        {"value": key, "count": tallies[key]}
        for key in sorted(tallies, key=lambda name: (-tallies[name], name))
    ]


def search_files(
    query: str,
    include_expired: bool = False,
    lifecycle: str | None = None,
    extension: str | None = None,
    page: int = 1,
    limit: int = 40,
) -> dict:
    cleaned = " ".join(query.split())
    page = max(1, page)
    empty = {
        "found": 0,
        "page": page,
        "groups": [],
        "facets": {"extension": [], "lifecycle": [], "organisation": []},
    }
    if not cleaned:
        return empty
    extensions = [
        piece.strip().lstrip(".").lower()
        for piece in (extension or "").split(",")
        if piece.strip()
    ]
    if lifecycle and lifecycle not in {"open", "closing_soon", "expired", "undated"}:
        lifecycle = None
    connection = client()
    ensure_collection(connection)
    parameters = {
        "q": cleaned,
        "query_by": "filename,content,title,organisation,tender_id",
        "query_by_weights": "4,3,2,1,2",
        "per_page": max(1, min(limit, 80)),
        "page": page,
        "highlight_full_fields": "content,filename,title",
        "highlight_affix_num_tokens": 8,
        "prefix": True,
        "num_typos": 2,
        "min_len_1typo": 3,
        "min_len_2typo": 5,
        "typo_tokens_threshold": 1,
        "drop_tokens_threshold": 1,
        "max_candidates": 250,
        "split_join_tokens": "fallback",
        "enable_synonyms": True,
        "synonym_num_typos": 1,
        "synonym_prefix": True,
        "facet_by": "extension",
        "max_facet_values": 20,
    }
    if _supports_infix(connection):
        # filename, content, title, organisation, tender_id
        parameters["infix"] = "fallback,off,fallback,fallback,off"
    filter_by = _build_filter(include_expired, lifecycle, extensions)
    if filter_by:
        parameters["filter_by"] = filter_by
    result = connection.collections[COLLECTION].documents.search(parameters)
    groups: dict[str, dict] = {}
    order: list[str] = []
    lifecycles: list[str] = []
    organisations: list[str] = []
    for hit in result.get("hits") or []:
        document = hit["document"]
        tender_id = document["tender_id"]
        closing = int(document.get("closing_at") or 0)
        stage = lifecycle_from_unix(closing)
        organisation = document.get("organisation") or "Unknown"
        lifecycles.append(stage)
        organisations.append(organisation)
        if tender_id not in groups:
            groups[tender_id] = {
                "tender_id": tender_id,
                "title": document.get("title") or "",
                "organisation": organisation,
                "lifecycle": stage,
                "closing_at": (
                    datetime.fromtimestamp(closing, timezone.utc).isoformat()
                    if closing
                    else None
                ),
                "files": [],
            }
            order.append(tender_id)
        relative = document.get("relative_path") or document.get("filename") or ""
        groups[tender_id]["files"].append(
            {
                "filename": document.get("filename") or relative,
                "relative_path": relative,
                "extension": document.get("extension") or "",
                "snippet": _snippet_html(hit),
                "score": hit.get("text_match") or hit.get("text_match_info", {}).get("score"),
                "matched_fields": _matched_fields(hit),
                "url": f"/files/{quote(tender_id)}/{quote(relative)}",
            }
        )
    return {
        "found": result.get("found") or 0,
        "page": page,
        "groups": [groups[tender_id] for tender_id in order],
        "facets": {
            "extension": _facet_rows(result, "extension"),
            "lifecycle": _count_map(lifecycles),
            "organisation": _count_map(organisations)[:12],
        },
    }
